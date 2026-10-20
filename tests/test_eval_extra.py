"""scripts/eval_extra.py: טבלת האמת של סט ההערכה הנוסף.
שורה לכל התחייבות (כמה שורות עם אותו מספר עמוד מקובצות), מרכזי מרובה בתא אחד מופרד ב-";".
בלי שום קריאה למודל: התוצאות נקראות מקובצי p<N>.json בתיקייה זמנית."""
import json

import pytest

from openpyxl import load_workbook

from scripts.eval_extra import (ALL_FIELDS, TRUTH_CANDIDATES, categories, cell_text, cells, commitment_cells,
                                load_truth, row_disagreements, truth_filled, truth_path)
from scripts.make_truth_xlsx import COLUMNS, TEXT_FORMAT, build_truth_xlsx
from src.schema import DateCandidate, Field_, NumberCandidate, PageExtraction, SenderMention

HEADER = "page,first_name,last_name,id,center,receipt_date,sender,commitment\n"
EMPTY_ROWS = "".join(f"{p},,,,,,,\n" for p in range(2, 7))


def write_truth(tmp_path, *rows: str):
    path = tmp_path / "truth.csv"
    path.write_text(HEADER + "".join(rows), encoding="utf-8")
    return path


def page_result(centers=("944963",), commitments=("107646",), first="ישראל", last="ישראלי", id_value="11223344",
                date="18.06.2026", sender="כללית"):
    """תוצאת מעבר ראשון סינתטית: כל השדות מודפסים ומתויגים, כדי שהכללים יחזירו ערך בלי דגלים."""
    numbers = [NumberCandidate(value=v, label="מרכזי", confidence=0.95, source="printed") for v in centers]
    numbers += [NumberCandidate(value=v, label="התחייבות", confidence=0.95, source="printed") for v in commitments]
    page = PageExtraction(
        is_commitment_form=True, page_quality="ok", rotation_degrees="0", numbers=numbers,
        id_numbers=[NumberCandidate(value=id_value, label='ת.ז', confidence=0.95, source="printed")],
        dates=[DateCandidate(value=date, kind="lab_stamp", confidence=0.95, source="printed", evidence="חותמת התקבל")],
        first_name=Field_(value=first, confidence=0.95, source="printed", label="שם פרטי"),
        last_name=Field_(value=last, confidence=0.95, source="printed", label="שם משפחה"),
        sender_mentions=[SenderMention(text=sender, where="header", role="sender", confidence=0.9, evidence="לוגו")],
    )
    return {"ok": True, "model": "fake", **page.model_dump()}


def save_page(out_dir, page: int, result: dict):
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / f"p{page}.json").write_text(
        json.dumps({"page": page, "pass1": result, "pass2": None, "second_pass": "skipped", "barcodes": None},
                   ensure_ascii=False), encoding="utf-8")


@pytest.fixture
def run_dir(tmp_path, monkeypatch):
    """תיקיית ריצה זמנית, כך שהבדיקה לא נוגעת ב-outputs/eval_extra."""
    out = tmp_path / "run"
    monkeypatch.setattr("scripts.eval_extra.OUT", out)
    return out


# ---- טבלת האמת: שורה לכל התחייבות ----

def test_two_rows_for_one_page_are_grouped(tmp_path):
    """שתי התחייבויות = שתי שורות עם אותו מספר עמוד: השנייה לא דורסת את הראשונה."""
    truth = load_truth(write_truth(tmp_path,
                                       "1,ישראל,ישראלי,11223344,944963,18.06.2026,כללית,107646\n",
                                       "1,ישראל,ישראלי,11223344,944963,18.06.2026,כללית,107647\n",
                                       EMPTY_ROWS))
    assert len(truth[1]) == 2
    assert [r["commitment"] for r in truth[1]] == ["107646", "107647"]
    for r in truth[1]:                              # פרטי המטופל חוזרים בשתי השורות
        assert (r["id"], r["first_name"], r["center"], r["date"], r["sender"]) == \
               ("11223344", "ישראל", "944963", "18.06.2026", "כללית")
    assert sorted(truth) == [1, 2, 3, 4, 5, 6]
    assert truth_filled(truth[1]) and not truth_filled(truth[2])


def test_multi_commitment_cell_also_expands(tmp_path):
    """תא התחייבות עם ";" נסבל גם הוא ומתפצל לשורה לכל התחייבות."""
    truth = load_truth(write_truth(tmp_path, "1,א,ב,11223344,944963,18.06.2026,כללית,107646 ; 107647\n", EMPTY_ROWS))
    assert [r["commitment"] for r in truth[1]] == ["107646", "107647"]


def test_row_disagreement_is_reported(tmp_path):
    """שדה שאמור לחזור זהה ונכתב שונה בין השורות: אזהרה (ההשוואה לוקחת את השורה הראשונה)."""
    truth = load_truth(write_truth(tmp_path,
                                       "1,ישראל,ישראלי,11223344,944963,18.06.2026,כללית,107646\n",
                                       "1,ישראל,ישראלי,99999999,944963,18.06.2026,מכבי,107647\n",
                                       EMPTY_ROWS))
    assert row_disagreements(truth[1]) == ["sender", "id"]
    assert row_disagreements(load_truth(write_truth(tmp_path, "1,א,ב,1,2,18.06.2026,כללית,7\n", EMPTY_ROWS))[1]) == []


# ---- מרכזי: כמה ערכים בתא אחד ----

def test_two_centers_in_one_cell_match_both_values(tmp_path, run_dir):
    """תא מרכזי "944963 ; 944964" מתפצל, ושני המספרים שנקראו מתאימים לו בדיוק."""
    truth = load_truth(write_truth(tmp_path, "1,ישראל,ישראלי,11223344,944963 ; 944964,18.06.2026,כללית,107646\n",
                                       EMPTY_ROWS))
    save_page(run_dir, 1, page_result(centers=("944963", "944964")))
    c = cells(truth, [1])[1]["center"]
    assert c["category"] == "confident_ok"
    assert c["truth"] == "944963 ; 944964"
    assert c["pred"] == "944963 ; 944964"


def test_two_centers_but_only_one_read_is_wrong(tmp_path, run_dir):
    truth = load_truth(write_truth(tmp_path, "1,ישראל,ישראלי,11223344,944963 ; 944964,18.06.2026,כללית,107646\n",
                                       EMPTY_ROWS))
    save_page(run_dir, 1, page_result(centers=("944963",)))
    c = cells(truth, [1])[1]["center"]
    assert c["category"] == "confident_wrong"
    assert c["pred"] == "944963"


def test_center_order_in_the_cell_does_not_matter(tmp_path, run_dir):
    truth = load_truth(write_truth(tmp_path, "1,ישראל,ישראלי,11223344,944964 ; 944963,18.06.2026,כללית,107646\n",
                                       EMPTY_ROWS))
    save_page(run_dir, 1, page_result(centers=("944963", "944964")))
    assert cells(truth, [1])[1]["center"]["category"] == "confident_ok"


# ---- התחייבות: תא לכל שורה ----

def pred_rows(*items: tuple[str | None, list[str]]):
    return [{"commitment": v, "commitment_flags": flags} for v, flags in items]


def test_each_commitment_row_is_compared_separately(tmp_path):
    truth = load_truth(write_truth(tmp_path, "1,א,ב,1,2,18.06.2026,כללית,107646\n",
                                       "1,א,ב,1,2,18.06.2026,כללית,107647\n", EMPTY_ROWS))
    out = commitment_cells(truth[1], pred_rows(("107646", []), ("107647", ["commitment_single_source"])))
    assert [c["category"] for c in out] == ["confident_ok", "orange_ok"]


def test_missing_second_commitment_is_counted_once(tmp_path):
    """המערכת הפיקה רק אחת משתי ההתחייבויות: אחת exact, השנייה ריקה ומסומנת."""
    truth = load_truth(write_truth(tmp_path, "1,א,ב,1,2,18.06.2026,כללית,107646\n",
                                       "1,א,ב,1,2,18.06.2026,כללית,107647\n", EMPTY_ROWS))
    out = commitment_cells(truth[1], pred_rows(("107646", []), (None, ["missing_required_field"])))
    assert [(c["truth"], c["pred"], c["category"]) for c in out] == \
           [("107646", "107646", "confident_ok"), ("107647", "", "empty_flagged")]


def test_wrong_commitment_without_flag_is_wrong_green(tmp_path):
    truth = load_truth(write_truth(tmp_path, "1,א,ב,1,2,18.06.2026,כללית,107646\n", EMPTY_ROWS))
    out = commitment_cells(truth[1], pred_rows(("999999", [])))
    assert [(c["truth"], c["pred"], c["category"]) for c in out] == [("107646", "999999", "confident_wrong")]


def test_extra_row_the_system_invented_is_wrong(tmp_path):
    truth = load_truth(write_truth(tmp_path, "1,א,ב,1,2,18.06.2026,כללית,107646\n", EMPTY_ROWS))
    out = commitment_cells(truth[1], pred_rows(("107646", []), ("999999", [])))
    assert [c["category"] for c in out] == ["confident_ok", "confident_wrong"]


def test_no_commitment_in_truth_and_none_found_is_one_empty_cell(tmp_path):
    truth = load_truth(write_truth(tmp_path, "1,א,ב,1,2,18.06.2026,כללית,\n", EMPTY_ROWS))
    out = commitment_cells(truth[1], pred_rows((None, ["no_commitment"])))
    assert [c["category"] for c in out] == ["empty_ok"]          # no_commitment לבדו אינו דגל לבדיקה


def test_commitment_counted_per_row_other_fields_per_page(tmp_path, run_dir):
    """עמוד עם שתי התחייבויות: n=2 להתחייבות, n=1 לשאר השדות (אותם פרטי מטופל בשתי השורות)."""
    truth = load_truth(write_truth(tmp_path, "1,ישראל,ישראלי,11223344,944963,18.06.2026,כללית,107646\n",
                                       "1,ישראל,ישראלי,11223344,944963,18.06.2026,כללית,107647\n", EMPTY_ROWS))
    save_page(run_dir, 1, page_result(commitments=("107646", "107647")))
    cmp = cells(truth, [1])
    assert sum(categories(cmp, "commitment", [1]).values()) == 2
    for f in ALL_FIELDS:
        if f != "commitment":
            assert sum(categories(cmp, f, [1]).values()) == 1
    assert categories(cmp, "commitment", [1])["confident_ok"] + \
           categories(cmp, "commitment", [1])["orange_ok"] == 2


# ---- טבלת אמת ב-xlsx ----

def write_xlsx(tmp_path, rows, columns=None, name="truth.xlsx", text_format=True):
    """xlsx בדיקה: שורת כותרת ואחריה השורות שניתנו. text_format=True: כל התאים בפורמט טקסט, כמו התבנית."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(columns or COLUMNS)
    for row in rows:
        ws.append(list(row))
    if text_format:
        for r in ws.iter_rows():
            for cell in r:
                cell.number_format = TEXT_FORMAT
    path = tmp_path / name
    wb.save(path)
    return path


def test_xlsx_keeps_leading_zero_in_id(tmp_path):
    """ת"ז שנכתבה כטקסט עם אפסים מובילים נקראת בדיוק כפי שנכתבה."""
    path = write_xlsx(tmp_path, [("1", "ישראל", "ישראלי", "00112233", "944963", "18.06.2026", "כללית", "0107646")])
    truth = load_truth(path)
    assert truth[1][0]["id"] == "00112233"
    assert truth[1][0]["commitment"] == "0107646"
    assert truth[1][0]["center"] == "944963"


def test_xlsx_numeric_cell_is_read_as_text_without_decimal_point(tmp_path):
    """תא שאקסל שמר כמספר (לא טקסט): נקרא כמחרוזת בלי ".0". האפסים המובילים אובדים באקסל עצמו,
    ולכן התבנית היא טקסט (scripts/make_truth_xlsx.py)."""
    path = write_xlsx(tmp_path, [(1, "א", "ב", 112233, 944963, "18.06.2026", "כללית", 107646)], text_format=False)
    truth = load_truth(path)
    assert truth[1][0]["id"] == "112233"
    assert truth[1][0]["commitment"] == "107646"
    assert sorted(truth) == [1]


def test_xlsx_real_date_cell_becomes_ddmmyyyy(tmp_path):
    """תא שאקסל המיר לתאריך אמיתי (ולכן בפורמט תאריך ולא טקסט) מוצג בפורמט הפלט של המערכת.
    בתא בפורמט טקסט תאריך נשאר מחרוזת כפי שהוקלדה, וזה המצב הרגיל בתבנית."""
    from datetime import datetime
    path = write_xlsx(tmp_path, [("1", "א", "ב", "112233", "944963", datetime(2026, 6, 18), "כללית", "107646")],
                      text_format=False)
    assert load_truth(path)[1][0]["date"] == "18.06.2026"
    as_text = write_xlsx(tmp_path, [("1", "א", "ב", "00112233", "944963", "18.06.2026", "כללית", "107646")],
                         name="text.xlsx")
    assert load_truth(as_text)[1][0]["date"] == "18.06.2026"


def test_xlsx_two_rows_for_one_page_are_grouped(tmp_path):
    path = write_xlsx(tmp_path, [("1", "ישראל", "ישראלי", "00112233", "944963 ; 944964", "18.06.2026", "כללית", "107646"),
                                 ("1", "ישראל", "ישראלי", "00112233", "944963 ; 944964", "18.06.2026", "כללית", "107647")])
    truth = load_truth(path)
    assert [r["commitment"] for r in truth[1]] == ["107646", "107647"]
    assert truth[1][0]["center"] == "944963 ; 944964"
    assert row_disagreements(truth[1]) == []


def test_xlsx_empty_rows_and_missing_column(tmp_path):
    empty = write_xlsx(tmp_path, [(str(p),) for p in range(1, 7)])
    truth = load_truth(empty)
    assert sorted(truth) == [1, 2, 3, 4, 5, 6]
    assert not any(truth_filled(truth[p]) for p in truth)
    bad = write_xlsx(tmp_path, [("1",)], columns=[c for c in COLUMNS if c != "id"], name="bad.xlsx")
    with pytest.raises(ValueError, match="id"):
        load_truth(bad)


def test_cell_text_variants():
    from datetime import date, datetime
    assert cell_text(None) is None
    assert cell_text("  ") is None
    assert cell_text(" 00112233 ") == "00112233"
    assert cell_text(107646) == "107646"
    assert cell_text(107646.0) == "107646"
    assert cell_text(date(2026, 6, 18)) == "18.06.2026"
    assert cell_text(datetime(2026, 6, 18, 7, 49)) == "18.06.2026"


# ---- התבנית שנוצרת, וסדר העדיפות של קובץ האמת ----

def test_build_truth_xlsx_is_text_formatted_and_readable(tmp_path):
    path = build_truth_xlsx(tmp_path / "truth.xlsx")
    ws = load_workbook(path).worksheets[0]
    assert [c.value for c in ws[1]] == COLUMNS
    assert ws.max_row == 7                                     # כותרת + 6 עמודים
    for row in ws.iter_rows():
        for cell in row:
            assert cell.number_format == TEXT_FORMAT           # כל התאים טקסט: אפסים מובילים נשמרים
    for i in range(1, len(COLUMNS) + 1):
        assert ws.column_dimensions[ws.cell(row=1, column=i).column_letter].number_format == TEXT_FORMAT
    truth = load_truth(path)
    assert sorted(truth) == [1, 2, 3, 4, 5, 6]
    assert all(len(truth[p]) == 1 and not truth_filled(truth[p]) for p in truth)


def test_truth_file_priority_xlsx_then_csv_then_template(tmp_path):
    """xlsx קודם ל-csv, ו-csv קודם לתבנית; --truth גובר על הכל."""
    xlsx, csv_file, template = tmp_path / "truth.xlsx", tmp_path / "truth.csv", tmp_path / "truth_template.csv"
    options = [xlsx, csv_file, template]
    assert truth_path([], options) == xlsx                     # אף אחד לא קיים: הראשון (ההודעה תאמר שחסר)
    template.write_text(HEADER, encoding="utf-8")
    assert truth_path([], options) == template
    csv_file.write_text(HEADER, encoding="utf-8")
    assert truth_path([], options) == csv_file
    build_truth_xlsx(xlsx)
    assert truth_path([], options) == xlsx
    assert truth_path([f"--truth={template}"], options) == template


def test_default_truth_candidates_order():
    assert [p.name for p in TRUTH_CANDIDATES] == ["truth.xlsx", "truth.csv", "truth_template.csv"]
