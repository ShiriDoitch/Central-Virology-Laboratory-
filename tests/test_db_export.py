"""ייצוא מהמסד (src/db_export.py, service.export_records): סינון, הגנת נוסחאות, קידוד, JSON. נתונים סינתטיים, בלי API."""
import csv
import io
import json

import pytest
from openpyxl import load_workbook

from src import db_export, repo, service
from tests.test_repo import fields, green_fields


@pytest.fixture
def conn():
    c = repo.connect(":memory:")
    uid = repo.create_upload(c, "טפסים.pdf", "sha", 3)
    repo.save_machine_record(c, uid, 1, 0, green_fields())
    repo.save_machine_record(c, uid, 2, 0, green_fields(first_name=("=HYPERLINK(\"x\")", "orange", ["handwritten"]), sender=("מכבי", "green"),
                                                       id=("000112233", "green")))
    repo.save_machine_record(c, uid, 3, 0, green_fields(sender=("-2+3", "green"), commitment=("@SUM(1)", "green")))
    yield c
    c.close()


def test_xlsx_respects_filter_and_protects_formulas(conn):
    data, mime, name = service.export_records(conn, service.RecordFilter(statuses=["pending", "ready"]), "xlsx")
    assert name == "records.xlsx" and "spreadsheetml" in mime
    ws = load_workbook(io.BytesIO(data)).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] == "שם פרטי" and len(rows) == 4          # כותרת + 3 רשומות
    body = [c for r in rows[1:] for c in r if isinstance(c, str)]
    assert "'=HYPERLINK(\"x\")" in body and "'-2+3" in body and "'@SUM(1)" in body
    assert not any(c.startswith(("=", "@")) for c in body)
    assert [r for r in rows[1:] if r[2] == "000112233"]         # אפס מוביל נשמר (טקסט)


def test_filter_limits_export_not_page_size(conn):
    data, _, _ = service.export_records(conn, service.RecordFilter(senders=["מכבי"]), "json")
    doc = json.loads(data)
    assert doc["count"] == 1 and doc["records"][0]["page"] == 2
    f = doc["records"][0]["fields"]["first_name"]
    assert f["flags"] == ["handwritten"] and f["color"] == "orange" and f["value"].startswith("=")      # JSON: הערך כפי שהוא


def test_csv_utf8_bom_and_protection(conn):
    data, mime, name = service.export_records(conn, service.RecordFilter(), "csv")
    assert data.startswith(b"\xef\xbb\xbf") and name == "records.csv"
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    assert rows[0][:2] == ["שם פרטי", "שם משפחה"] and len(rows) == 4
    flat = [c for r in rows[1:] for c in r]
    assert "'-2+3" in flat and "'@SUM(1)" in flat and "טפסים.pdf" in flat


def test_edited_and_approved_reflected_in_all_formats(conn):
    rid = repo.search_records(conn, search="@SUM", limit=1)[0][0]["id"]
    repo.edit_field(conn, rid, "sender", "כללית")
    repo.approve_record(conn, rid)
    doc = json.loads(service.export_records(conn, service.RecordFilter(statuses=["approved"]), "json")[0])
    assert doc["count"] == 1 and doc["records"][0]["status"] == "approved"
    assert doc["records"][0]["fields"]["sender"]["human_edited"] and doc["records"][0]["fields"]["sender"]["machine_value"] == "-2+3"
    ws = load_workbook(io.BytesIO(service.export_records(conn, service.RecordFilter(statuses=["approved"]), "xlsx")[0])).active
    assert ws.cell(2, 6).value == "כללית" and "תוקן ידנית" in ws.cell(2, 6).comment.text


def test_empty_selection_and_bad_format(conn):
    data, _, _ = service.export_records(conn, service.RecordFilter(search="nothing-matches"), "json")
    assert json.loads(data)["count"] == 0
    with pytest.raises(ValueError):
        service.export_records(conn, service.RecordFilter(), "pdf")


# ---------- סימון "לא בטוח" ב-CSV (21.10.2026) ----------

def test_csv_marks_the_uncertain_fields(conn):
    """ל-CSV אין צבעים, ולכן הוא נושא עמודה שמונה את השדות שדורשים בדיקה (כתום או צהוב).
    בלעדיה היה זה הייצוא היחיד שבו "לא בטוח" נעלם לגמרי."""
    data, _, _ = service.export_records(conn, service.RecordFilter(), "csv")
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    assert rows[0][-1] == db_export.UNCERTAIN_COLUMN
    by_page = {r[9]: r[-1] for r in rows[1:]}            # עמודה "עמוד"
    assert by_page["2"] == "שם פרטי"                      # היחיד הכתום באותה רשומה
    assert by_page["1"] == "" and by_page["3"] == ""      # רשומות ירוקות לגמרי


def test_csv_uncertain_column_counts_yellow_and_flagged_empty_but_not_confirmed(conn):
    uid = repo.create_upload(conn, "ב.pdf", "sha-b", 1)
    rid, _ = repo.save_machine_record(conn, uid, 9, 0, green_fields(
        first_name=("ישראל", "yellow", ["handwritten"]), last_name=(None, "orange", ["handwritten_unconfirmed"]),
        sender=("כללית", "green")))
    repo.edit_field(conn, rid, "center", "944963")        # תיקון ידני: confirmed, לא נספר
    data, _, _ = service.export_records(conn, service.RecordFilter(upload_id=uid), "csv")
    row = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))[1]
    assert row[-1] == "שם פרטי, שם משפחה"
