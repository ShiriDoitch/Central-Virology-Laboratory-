import json

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from src import extractor, rules
from src.cache import load_snapshot
from src.export import export_excel
from src.rules import F_FATHER_NAME, decide_name
from src.schema import DateCandidate, Field_, PageExtraction, SenderMention
from src.senders import F_PATIENT_FIELD_ONLY, F_UNKNOWN_SENDER, decide_sender, load_matcher

M = load_matcher()


def mention(text, where="header", role="sender", conf=1.0):
    return SenderMention(text=text, where=where, role=role, confidence=conf)


# ---------- סכמה ופרומפט ----------

def test_schema_v2_fields():
    page = PageExtraction(is_commitment_form=True, page_quality="ok")
    assert page.rotation_degrees == "0"
    assert Field_().label is None
    assert SenderMention(text="x", where="patient_field", role="unknown").evidence is None
    with pytest.raises(ValidationError):
        DateCandidate(value="1.1.2026", kind="lab_received")                  # הסוג הישן כבר לא תקף
    with pytest.raises(ValidationError):
        PageExtraction(is_commitment_form=True, page_quality="ok", rotation_degrees="45")
    DateCandidate(value="1.1.2026", kind="lab_stamp")
    DateCandidate(value="1.1.2026", kind="handwritten_date")


def test_prompt_v2_mentions_the_new_concepts_and_not_the_old_ones():
    p = extractor.PROMPT
    for phrase in ("lab_stamp", "handwritten_date", "patient_field", "rotation_degrees", "evidence",
                   "father", "not legible", "SENDER", "label"):
        assert phrase in p, phrase
    assert "lab_received" not in p and "rotation_needed" not in p
    assert "yellow" not in p.lower() and "צהוב" not in p                         # הוראות בלי צבע
    assert "label" in extractor.PROMPT_REGIONS and "father" in extractor.PROMPT_REGIONS.lower()


# ---------- תצלום ישן ----------

def test_legacy_snapshot_is_converted_and_validates(tmp_path):
    legacy_page = {
        "ok": True, "model": "m", "is_commitment_form": True, "page_quality": "ok", "rotation_needed": 270,
        "first_name": {"value": None, "confidence": 0.0, "evidence": None, "source": "unknown"},
        "last_name": {"value": None, "confidence": 0.0, "evidence": None, "source": "unknown"},
        "id_numbers": [], "numbers": [], "sender_mentions": [],
        "dates": [
            {"value": "05-08-2026", "kind": "lab_received", "evidence": "קבלה", "confidence": 0.9, "source": "printed"},
            {"value": "12/08/2026", "kind": "lab_received", "evidence": "תאריך קבלה", "confidence": 1.0, "source": "printed"},
            {"value": "16/8/26", "kind": "lab_received", "evidence": None, "confidence": 0.9, "source": "handwritten"},
            {"value": "4-8-26", "kind": "other", "evidence": "תאריך", "confidence": 0.9, "source": "handwritten"},
        ],
    }
    path = tmp_path / "snap.json"
    path.write_text(json.dumps({"pages": {"14": legacy_page}}, ensure_ascii=False), encoding="utf-8")
    page = load_snapshot(path)[14]
    assert [d["kind"] for d in page["dates"]] == ["lab_stamp", "receipt_date", "handwritten_date", "other"]
    model = PageExtraction.model_validate(page)                                  # עובר אימות
    assert model.rotation_degrees == "0" and model.first_name.label is None
    d = rules.decide_receipt_date(model.dates)
    assert d.value == "05.08.2026" and d.alternatives == ["16.08.2026", "04.08.2026"]   # חותמת מנצחת, כתב יד בהערה


def test_real_snapshot_validates_if_present():
    snap = load_snapshot()
    for result in snap.values():
        PageExtraction.model_validate(result)


# ---------- patient_field ----------

def test_patient_field_hmo_does_not_beat_institution_header():
    """עמוד 26 בגרסה v2: מאוחדת היא patient_field ולכן לא מתחרה בכותרת המוסד."""
    ms = [mention("המרכז הרפואי שמיר (אסף הרופא) אגף המעבדות", "header", "destination", 0.99),
          mention("מאוחדת", "patient_field", "unknown", 0.95)]
    d = decide_sender(ms, M)
    assert d.value == "המרכז הרפואי שמיר" and d.flags == []
    assert d.alternatives == ["מאוחדת"]


def test_patient_field_does_not_vote_even_with_only_a_stamp_present():
    ms = [mention("מכבי", "patient_field", "sender"), mention("הלל יפה", "stamp", "sender")]
    d = decide_sender(ms, M)
    assert d.value == "מרכז רפואי הלל יפה" and d.alternatives == ["מכבי"]


def test_patient_field_only_is_unflagged_and_keeps_a_hint():
    d = decide_sender([mention("כללית", "patient_field", "unknown")], M)
    assert d.value == "כללית" and d.flags == [] and d.hints == [F_PATIENT_FIELD_ONLY] and not d.uncertain


def test_patient_field_only_does_not_hide_other_flags_or_unknown_senders():
    d = decide_sender([mention("חברה לא מוכרת בע\"מ", "header", "sender"), mention("כללית", "patient_field", "unknown")], M)
    assert d.flags and d.value == "אחר"                                   # גורם לא מוכר נשאר כתום
    d = decide_sender([mention("כללית", "patient_field", "unknown"), mention("מכבי", "patient_field", "unknown")], M)
    assert d.value == "כללית" and d.alternatives == ["מכבי"]


def test_unknown_logo_beats_patient_field_and_keeps_it_in_alternatives():
    ms = [mention("מכון מור", "logo", "sender"), mention("מאוחדת", "patient_field", "unknown")]
    d = decide_sender(ms, M)
    assert d.value == "אחר" and d.flags == [F_UNKNOWN_SENDER] and d.alternatives == ["מכון מור", "מאוחדת"]


def test_patient_field_equal_to_winner_is_not_listed_twice():
    ms = [mention("מאוחדת", "logo", "sender"), mention("מאוחדת", "patient_field", "unknown")]
    d = decide_sender(ms, M)
    assert d.value == "מאוחדת" and d.alternatives == [] and d.flags == []


def test_lab_name_as_patient_field_is_still_excluded():
    assert decide_sender([mention("המרכז הרפואי המשולב", "patient_field", "unknown")], M).value is None


def test_sender_evidence_field_is_accepted_and_ignored_by_decision():
    ms = [SenderMention(text="מכבי", where="logo", role="sender", evidence="לוגו מכבי בראש הטופס", confidence=1.0)]
    assert decide_sender(ms, M).value == "מכבי"


# ---------- תווית שם האב ----------

def test_father_label_without_evidence_is_rejected():
    d = decide_name(Field_(value="אריה", confidence=0.9, source="handwritten", label="שם האב"))
    assert d.value is None and d.flags == [F_FATHER_NAME] and d.alternatives == ["אריה"]


def test_patient_label_is_kept_as_before():
    d = decide_name(Field_(value="דנה", confidence=0.9, source="printed", label="שם המטופל"))
    assert d.value == "דנה" and d.flags == []


@pytest.mark.parametrize("label", ["שם פרטי", "שם משפחה"])
def test_non_father_labels_do_not_trigger(label):
    assert decide_name(Field_(value="דנה", confidence=0.9, source="printed", label=label)).value == "דנה"


@pytest.mark.parametrize("label", ["אב-טיפוס", "תל אביב"])
def test_non_father_unknown_labels_are_not_father_but_unknown(label):
    d = decide_name(Field_(value="דנה", confidence=0.9, source="printed", label=label))
    assert F_FATHER_NAME not in d.flags and d.flags == ["name_label_unknown"] and d.value is None


# ---------- אקסל: סיבוב ודגלי עמוד ----------

def test_export_shows_rotation_and_page_flags(tmp_path):
    p = PageExtraction(is_commitment_form=True, page_quality="ok")
    page = {"page": 14, "pass1": {"ok": True, "model": "m", **p.model_dump()}, "pass2": None,
            "rotation": {"local_axis": "vertical", "reported": "270", "applied": None, "attempts": 1, "unresolved": True},
            "page_flags": ["rotation_unresolved"]}
    path = tmp_path / "x.xlsx"
    summary = export_excel([page], path)
    audit = load_workbook(path)["audit"]
    headers = [c.value for c in audit[1]]
    row = {h: c.value for h, c in zip(headers, next(audit.iter_rows(min_row=2)))}
    assert "vertical" in row["סיבוב"] and "270" in row["סיבוב"] and "לא סובב" in row["סיבוב"]
    assert row["דגלי עמוד"] == "rotation_unresolved"
    assert summary["flags"]["rotation_unresolved"] == 1
