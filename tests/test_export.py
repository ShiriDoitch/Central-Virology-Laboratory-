import pytest
from openpyxl import load_workbook

from src import export
from src.export import COLUMNS, ORANGE, build_rows, export_excel, protect, skip_reason
from src.schema import DateCandidate, Field_, NumberCandidate, PageExtraction, RegionsExtraction, SenderMention

ORANGE_RGB = ORANGE.fgColor.rgb[-6:]
GREEN_RGB = export.GREEN.fgColor.rgb[-6:]
YELLOW_RGB = export.YELLOW.fgColor.rgb[-6:]


def n(value, label=None, source="printed", conf=1.0):
    return NumberCandidate(value=value, label=label, source=source, confidence=conf)


def make_page(page_no, is_form=True, **kw):
    p = PageExtraction(is_commitment_form=is_form, page_quality="ok", **kw)
    return {"page": page_no, "pass1": {"ok": True, "model": "model-a", **p.model_dump()}, "pass2": None}


CONFIRMING_PASS2 = {"ok": True, "model": "m", **RegionsExtraction.model_validate({
    "patient": {"region_ok": True},
    "barcode": {"region_ok": True, "central_number": {"value": "939376", "label": "מרכזי", "confidence": 0.95, "source": "printed"}},
    "commitment": {"region_ok": True, "number": {"value": "730578626", "confidence": 0.95, "source": "printed"}},
    "receipt": {"region_ok": True, "date": {"value": "18.06.2026", "confidence": 0.95, "source": "printed"}},    # קריאה שנייה של התאריך
}).model_dump()}


def good_page(page_no=1, confirmed=True, **extra):
    kw = dict(
        first_name=Field_(value="דנה", confidence=0.95, source="printed", label="שם פרטי"),
        last_name=Field_(value="לוי", confidence=0.95, source="printed", label="שם משפחה"),
        id_numbers=[n("11223344", "ת.ז.")],
        dates=[DateCandidate(value="18-06-2026", kind="lab_stamp", evidence="stamp", source="printed", confidence=1.0)],
        numbers=[n("939376", "מרכזי"), n("730578626", "התחייבות מספר")],
        sender_mentions=[SenderMention(text="מאוחדת", where="logo", role="sender", confidence=1.0)],
    )
    kw.update(extra)
    page = make_page(page_no, **kw)
    # המעבר השני מאשר את המרכזי ואת מספר ההתחייבות (רק בעמוד ברירת המחדל): ירוק
    page["pass2"] = CONFIRMING_PASS2 if confirmed and "numbers" not in extra else None
    return page


def with_second_pass_id(page, digits="11223344"):
    """המעבר השני (חיתוך תקין) קרא את ת"ז בכתב יד: שתי קריאות עצמאיות."""
    import copy
    page = copy.deepcopy(page)
    pass2 = copy.deepcopy(page["pass2"] or CONFIRMING_PASS2)
    pass2["patient"]["region_ok"] = True
    pass2["patient"]["id_number"] = {"value": digits, "confidence": 0.9, "evidence": None, "label": "ת.ז.", "source": "handwritten",
                                     "alternatives": []}
    page["pass2"] = pass2
    return page


def with_second_pass_name(page, value="ישראל"):
    """המעבר השני (חיתוך תקין) קרא את השם הפרטי בכתב יד: שתי קריאות עצמאיות."""
    import copy
    page = copy.deepcopy(page)
    pass2 = copy.deepcopy(page["pass2"] or CONFIRMING_PASS2)
    pass2["patient"]["region_ok"] = True
    pass2["patient"]["first_name"] = {"value": value, "confidence": 0.9, "evidence": None, "label": "שם פרטי",
                                      "source": "handwritten", "alternatives": []}
    page["pass2"] = pass2
    return page


def handwritten_first_name_page(page_no=1):
    return with_second_pass_name(good_page(page_no, first_name=Field_(value="ישראל", confidence=0.99, source="handwritten",
                                                                      label="שם פרטי")))


def save_and_load(tmp_path, pages):
    path = tmp_path / "out.xlsx"
    summary = export_excel(pages, path)
    return load_workbook(path), summary


def main_cells(wb, row=2):
    ws = wb["תוצאות"]
    return [ws.cell(row=row, column=c) for c in range(1, len(COLUMNS) + 1)]


# ---------- מבנה ותוכן ----------

def test_main_sheet_headers_and_clean_row(tmp_path):
    wb, _ = save_and_load(tmp_path, [good_page()])
    ws = wb["תוצאות"]
    assert [c.value for c in ws[1]] == [t for _, t in COLUMNS]
    assert [c.value for c in main_cells(wb)] == ["דנה", "לוי", "11223344", "939376", "18.06.2026", "מאוחדת", "730578626"]
    assert all(c.fill.fgColor.rgb[-6:] != ORANGE_RGB for c in main_cells(wb))      # הכול ירוק
    assert wb.sheetnames == ["תוצאות", "audit", "סיכום", "דולגו"]


def rgb(cell):
    return cell.fill.fgColor.rgb[-6:] if cell.fill and cell.fill.fill_type == "solid" else None


def test_colors_green_for_value_without_flag(tmp_path):
    wb, _ = save_and_load(tmp_path, [good_page()])
    assert all(rgb(c) == GREEN_RGB for c in main_cells(wb))            # כל השורה: ערך בלי דגל


def test_colors_orange_for_flag_and_white_for_empty_without_flag(tmp_path):
    page = good_page(id_numbers=[n("11223344", "ת.ז.", source="handwritten", conf=0.3)],   # כתב יד + ביטחון נמוך: כתום
                     sender_mentions=[])                                              # ריק בלי דגל: לבן
    wb, _ = save_and_load(tmp_path, [page])
    cells = dict(zip([k for k, _ in COLUMNS], main_cells(wb)))
    assert rgb(cells["id"]) == ORANGE_RGB
    assert cells["sender"].value is None and rgb(cells["sender"]) is None          # בלי מילוי בכלל
    assert rgb(cells["first_name"]) == GREEN_RGB


def test_empty_name_is_always_flagged_for_review(tmp_path):
    """שם ריק אינו ריק בשקט (20.10.2026): בלי דגל אחר שמסביר אותו הוא מקבל missing_required_field וכתום."""
    wb, _ = save_and_load(tmp_path, [good_page(last_name=Field_())])
    cells = dict(zip([k for k, _ in COLUMNS], main_cells(wb)))
    assert cells["last_name"].value is None and rgb(cells["last_name"]) == ORANGE_RGB
    assert "missing_required_field" in cells["last_name"].comment.text


def test_empty_flagged_cell_is_orange_not_white_or_green(tmp_path):
    page = good_page(id_numbers=[n("00112233", "ת.ז."), n("11223344", "ת.ז.")])    # סתירה: ריק מסומן
    wb, _ = save_and_load(tmp_path, [page])
    assert rgb(main_cells(wb)[2]) == ORANGE_RGB


def test_summary_sheet_has_color_legend(tmp_path):
    wb, _ = save_and_load(tmp_path, [good_page()])
    ws = wb["סיכום"]
    legend = {rgb(ws.cell(row=r, column=1)): ws.cell(row=r, column=2).value for r in range(2, 5)}
    assert set(legend) == {GREEN_RGB, ORANGE_RGB, None}
    assert "ירוק" in legend[GREEN_RGB] and "כתום" in legend[ORANGE_RGB] and "לבן" in legend[None]
    assert ws["A1"].value == "מקרא (גיליון תוצאות)"


def test_id_keeps_leading_zeros_as_text(tmp_path):
    page = good_page(id_numbers=[n("0011223344", "ת.ז.")])
    wb, _ = save_and_load(tmp_path, [page])
    cell = main_cells(wb)[2]
    assert cell.value == "0011223344" and cell.data_type == "s"


# ---------- הגנה מנוסחאות ----------

@pytest.mark.parametrize("value", ["=1+1", "+1", "-2", "@SUM(A1)", "=HYPERLINK(\"http://x\")"])
def test_formula_injection_gets_apostrophe(tmp_path, value):
    page = good_page(first_name=Field_(value=value, confidence=0.95, source="printed", label="שם פרטי"))
    wb, _ = save_and_load(tmp_path, [page])
    cell = main_cells(wb)[0]
    assert cell.value == "'" + value
    assert cell.data_type == "s"            # לא נוסחה
    audit = wb["audit"]
    assert any(c.value == "'" + value for row in audit.iter_rows(min_row=2) for c in row)    # גם ב-audit


def test_protect_unit():
    assert protect("=A1") == "'=A1"
    assert protect("+1") == "'+1" and protect("-1") == "'-1" and protect("@x") == "'@x"
    assert protect("\t=A1") == "'\t=A1"
    assert protect("ישראל") == "ישראל" and protect("18.06.2026") == "18.06.2026"
    assert protect(None) is None and protect(5) == 5


# ---------- תא כתום ----------

def test_handwritten_name_read_once_is_shown_in_orange(tmp_path):
    """שונה ב-20.10.2026: קריאה אחת בכתב יד עם מעבר שני תקין מוצגת בתא ככתום (הצעה לאישור), לא ריקה."""
    page = good_page(first_name=Field_(value="דנה", confidence=0.95, source="handwritten"))
    wb, _ = save_and_load(tmp_path, [page])
    first, last = main_cells(wb)[0], main_cells(wb)[1]
    assert first.value == "דנה" and first.fill.fgColor.rgb[-6:] == ORANGE_RGB
    assert "handwritten_unconfirmed" in first.comment.text
    assert not first.comment.text.startswith(export.NOT_READ_NOTE)        # יש ערך בתא, אז אין "לא נקרא"
    assert last.fill.fgColor.rgb[-6:] != ORANGE_RGB


def test_name_without_a_second_pass_stays_empty_with_the_type_it_note(tmp_path):
    """בלי מעבר שני תקין הקריאה הבודדת נשארת רעש: תא ריק וכתום, "לא נקרא, יש להקליד"."""
    page = good_page(first_name=Field_(value="דנה", confidence=0.95, source="handwritten"), confirmed=False)
    wb, _ = save_and_load(tmp_path, [page])
    first = main_cells(wb)[0]
    assert first.value is None and first.fill.fgColor.rgb[-6:] == ORANGE_RGB
    assert first.comment.text.startswith(export.NOT_READ_NOTE)
    assert "דנה" not in first.comment.text


def test_handwritten_name_with_two_agreeing_readings_is_yellow_with_a_glance_comment(tmp_path):
    wb, _ = save_and_load(tmp_path, [handwritten_first_name_page()])
    cell = main_cells(wb)[0]
    assert cell.value == "ישראל" and cell.fill.fgColor.rgb[-6:] == YELLOW_RGB      # לא כתום ולא ירוק
    assert cell.comment.text.startswith(export.YELLOW_NOTE) and "מבט" in cell.comment.text


def test_id_passing_checksum_with_two_agreeing_readings_is_green_even_handwritten(tmp_path):
    page = with_second_pass_id(good_page(id_numbers=[n("11223344", "ת.ז.", source="handwritten")]))
    wb, _ = save_and_load(tmp_path, [page])
    cell = main_cells(wb)[2]
    assert cell.value == "11223344" and cell.fill.fgColor.rgb[-6:] == GREEN_RGB


def test_id_failing_checksum_is_orange_even_with_two_agreeing_readings(tmp_path):
    page = with_second_pass_id(good_page(id_numbers=[n("11923344", "ת.ז.", source="handwritten")]), digits="11923344")
    wb, _ = save_and_load(tmp_path, [page])
    cell = main_cells(wb)[2]
    assert cell.value == "11923344" and cell.fill.fgColor.rgb[-6:] == ORANGE_RGB       # הערך לא משתנה
    assert "id_checksum_failed" in cell.comment.text


def test_single_reading_handwritten_id_and_any_handwritten_date_are_orange_not_yellow(tmp_path):
    """"שתי קריאות מסכימות" נכון רק למי שנקרא פעמיים (ת"ז ושמות). תאריך נקרא במעבר הראשון בלבד: תמיד כתום."""
    single = good_page(id_numbers=[n("11223344", "ת.ז.", source="handwritten")])               # בלי ת"ז במעבר שני
    date = DateCandidate(value="10.08.2026", kind="handwritten_date", evidence="handwriting", source="handwritten", confidence=0.95)
    dated = good_page(2, dates=[date])
    dated["pass2"] = {**dated["pass2"], "receipt": {**dated["pass2"]["receipt"], "region_ok": False}}   # בלי קריאה שנייה של התאריך
    wb, _ = save_and_load(tmp_path, [single, dated])
    assert rgb(main_cells(wb, 2)[2]) == ORANGE_RGB                                              # ת"ז בקריאה אחת
    assert main_cells(wb, 3)[4].value == "10.08.2026" and rgb(main_cells(wb, 3)[4]) == ORANGE_RGB   # תאריך בכתב יד
    assert "קריאות מסכימות" not in main_cells(wb, 3)[4].comment.text


def test_value_with_handwritten_and_another_flag_stays_orange(tmp_path):
    page = good_page(id_numbers=[n("11223344", "ת.ז.", source="handwritten", conf=0.3)])      # + low_confidence
    cell = main_cells(save_and_load(tmp_path, [page])[0])[2]
    assert cell.fill.fgColor.rgb[-6:] == ORANGE_RGB and "handwritten" in cell.comment.text


# ---------- תא ריק מסומן ----------

def test_empty_flagged_cell_shows_raw_in_orange(tmp_path):
    bad_date = DateCandidate(value="31.06.2026", kind="lab_stamp", evidence="stamp", source="printed", confidence=1.0)
    wb, _ = save_and_load(tmp_path, [good_page(dates=[bad_date])])
    cell = main_cells(wb)[4]
    assert cell.value == "31.06.2026"                                # ה-raw
    assert cell.fill.fgColor.rgb[-6:] == ORANGE_RGB
    assert "invalid_date" in cell.comment.text and "raw: 31.06.2026" in cell.comment.text


def test_empty_flagged_without_raw_is_blank_orange_with_alternatives(tmp_path):
    page = good_page(id_numbers=[n("00112233", "ת.ז."), n("11223344", "ת.ז.")])    # סתירה
    wb, _ = save_and_load(tmp_path, [page])
    cell = main_cells(wb)[2]
    assert cell.value is None
    assert cell.fill.fgColor.rgb[-6:] == ORANGE_RGB
    assert "conflict" in cell.comment.text and "00112233" in cell.comment.text and "11223344" in cell.comment.text


def test_empty_unflagged_is_blank_not_orange(tmp_path):
    wb, _ = save_and_load(tmp_path, [good_page(sender_mentions=[])])
    cell = dict(zip([k for k, _ in COLUMNS], main_cells(wb)))["sender"]
    assert cell.value is None and cell.fill.fgColor.rgb[-6:] != ORANGE_RGB


# ---------- עמוד שאינו טופס ----------

def test_non_form_page_goes_to_skipped_sheet(tmp_path):
    pages = [good_page(1), make_page(2, is_form=False),
             {"page": 3, "pass1": {"ok": False, "error": "503"}, "pass2": None}]
    wb, summary = save_and_load(tmp_path, pages)
    assert wb["תוצאות"].max_row == 2                                  # כותרת + שורה אחת
    skipped = {row[0].value: row[1].value for row in wb["דולגו"].iter_rows(min_row=2)}
    assert set(skipped) == {2, 3}
    assert "לא טופס התחייבות" in skipped[2] and "נכשל" in skipped[3]
    assert summary["skipped"] == 2 and summary["rows"] == 1


def test_skip_reason_none_for_form():
    assert skip_reason(good_page()) is None
    assert skip_reason({"page": 1, "pass1": None}) is not None


def test_form_marked_not_commitment_but_with_number_is_kept(tmp_path):
    page = make_page(1, is_form=False, numbers=[n("483263113", "מס' הדגימה")])
    wb, summary = save_and_load(tmp_path, [page])
    assert summary["rows"] == 1 and main_cells(wb)[6].value == "483263113"
    assert main_cells(wb)[6].fill.fgColor.rgb[-6:] == ORANGE_RGB       # model_says_not_commitment_form


# ---------- ריבוי מרכזיים והתחייבויות ----------

def test_multiple_centers_joined_with_semicolon(tmp_path):
    page = good_page(numbers=[n("937309", "מרכזי"), n("937310", "מרכזי"), n("937311", "מרכזי"),
                              n("50364186", "דרישה")])
    wb, _ = save_and_load(tmp_path, [page])
    assert main_cells(wb)[3].value == "937309 ; 937310 ; 937311"


def test_one_row_per_commitment_with_same_patient_details(tmp_path):
    page = good_page(numbers=[n("948062", "מרכזי"), n("260005102632", "מספר התחייבות"),
                              n("260005102611", "מספר התחייבות"), n("260005102644", "מספר התחייבות")])
    wb, summary = save_and_load(tmp_path, [page])
    ws = wb["תוצאות"]
    assert summary["rows"] == 3 and ws.max_row == 4
    assert {ws.cell(row=r, column=7).value for r in (2, 3, 4)} == {"260005102632", "260005102611", "260005102644"}
    assert {ws.cell(row=r, column=4).value for r in (2, 3, 4)} == {"948062"}


def test_form_without_commitment_is_orange_empty(tmp_path):
    page = good_page(numbers=[n("939376", "מרכזי")])
    wb, _ = save_and_load(tmp_path, [page])
    cell = main_cells(wb)[6]
    assert cell.value is None and cell.fill.fgColor.rgb[-6:] == ORANGE_RGB


# ---------- audit וסיכום ----------

def test_audit_has_row_per_field_with_page_model_and_readings(tmp_path):
    wb, _ = save_and_load(tmp_path, [good_page(7, confirmed=False)])
    audit = wb["audit"]
    assert audit.max_row == 1 + len(COLUMNS)
    headers = [c.value for c in audit[1]]
    rows = [{h: c.value for h, c in zip(headers, row)} for row in audit.iter_rows(min_row=2)]
    by_field = {r["שדה"]: r for r in rows}
    assert set(by_field) == {k for k, _ in COLUMNS}
    r = by_field["id"]
    assert r["עמוד"] == 7 and r["מעבר ראשון: מודל"] == "model-a" and r["ערך סופי"] == "11223344"
    assert "11223344" in r["מעבר ראשון: קריאות"] and r["מעבר שני: סטטוס"] == "לא רץ"
    assert "stamp" in by_field["date"]["מעבר ראשון: קריאות"] and "lab_stamp" in by_field["date"]["מעבר ראשון: קריאות"]


def test_audit_includes_second_pass_readings_and_status(tmp_path):
    page = good_page(3)
    page["pass2"] = {
        "ok": True, "model": "strong-model",
        "patient": {"region_ok": True, "id_number": {"value": "11223344", "confidence": 0.9, "source": "handwritten",
                                                       "evidence": None, "alternatives": ["11223341"]},
                    "first_name": {"value": "דנה", "confidence": 0.9, "source": "handwritten", "evidence": None, "alternatives": []},
                    "last_name": {"value": None, "confidence": 0.0, "source": "unknown", "evidence": None, "alternatives": []},
                    "birth_date": {"value": None, "confidence": 0.0, "source": "unknown", "evidence": None, "alternatives": []}},
        "barcode": {"region_ok": True, "visible_words": [], "numbers": [],
                    "central_number": {"value": None, "confidence": 0.0, "source": "unknown", "evidence": None, "alternatives": []}},
    }
    wb, _ = save_and_load(tmp_path, [page])
    audit = wb["audit"]
    headers = [c.value for c in audit[1]]
    rows = [{h: c.value for h, c in zip(headers, row)} for row in audit.iter_rows(min_row=2)]
    r = next(x for x in rows if x["שדה"] == "id")
    assert r["מעבר שני: סטטוס"] == "רץ" and r["מעבר שני: מודל"] == "strong-model"
    assert "11223341" in r["מעבר שני: קריאות"]


def test_summary_counts(tmp_path):
    pages = [
        good_page(1),                                                                   # הכול ירוק
        good_page(2, id_numbers=[n("11223344", "ת.ז.", source="handwritten")]),         # כתום אחד (ת"ז בכתב יד, קריאה אחת)
        good_page(3, id_numbers=[n("00112233", "ת.ז."), n("11223344", "ת.ז.")]),        # ריק מסומן אחד
        good_page(4, sender_mentions=[]),                                               # ריק אחד
        make_page(5, is_form=False),
    ]
    pages[0]["stats"] = {"pass1": {"api_calls": 2, "quota_failures": 1}, "pass2": {}}
    _, s = save_and_load(tmp_path, pages)
    assert s["rows"] == 4 and s["skipped"] == 1 and s["fields_total"] == 28
    assert s["totals"] == {"green": 25, "orange": 1, "empty_flagged": 1, "empty": 1}
    assert s["by_field"]["id"]["orange"] == 1 and s["by_field"]["id"]["empty_flagged"] == 1
    assert s["flags"]["handwritten"] == 1 and s["flags"]["conflict"] == 1
    assert s["api"]["api_calls"] == 2


def test_summary_counts_unconfirmed_handwritten_names_as_orange(tmp_path):
    """שונה ב-20.10.2026: ההצעה מוצגת בתא, ולכן היא נספרת ככתום ולא כ"ריק מסומן"."""
    page = good_page(1, first_name=Field_(value="דנה", confidence=0.9, source="handwritten"))
    _, s = save_and_load(tmp_path, [page])
    assert s["by_field"]["first_name"] == {"orange": 1}
    assert s["flags"] == {"handwritten_unconfirmed": 1}


def test_summary_counts_a_name_without_a_second_pass_as_empty_flagged(tmp_path):
    page = good_page(1, first_name=Field_(value="דנה", confidence=0.9, source="handwritten"), confirmed=False)
    _, s = save_and_load(tmp_path, [page])
    assert s["by_field"]["first_name"] == {"empty_flagged": 1}


def test_summary_sheet_legend_and_counts_include_yellow(tmp_path):
    wb, s = save_and_load(tmp_path, [handwritten_first_name_page(1)])
    ws = wb["סיכום"]
    legend = {rgb(ws.cell(row=r, column=1)): ws.cell(row=r, column=2).value for r in range(2, 7)}
    assert "צהוב" in legend[YELLOW_RGB] and "מבט" in legend[YELLOW_RGB] and 'ת"ז או שם' in legend[YELLOW_RGB]
    texts = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    assert 'שדות צהובים (ת"ז / שם בכתב יד, שתי קריאות מסכימות)' in texts and "צהוב" in texts
    assert s["by_field"]["first_name"] == {"yellow": 1} and s["totals"]["yellow"] == 1


def test_summary_sheet_is_written(tmp_path):
    wb, _ = save_and_load(tmp_path, [good_page(1)])
    texts = [c.value for row in wb["סיכום"].iter_rows() for c in row if c.value is not None]
    assert "שורות (התחייבויות)" in texts and "שכיחות דגלים" in texts


def test_audit_reports_daily_quota_exhausted(tmp_path):
    page = good_page(2)
    page.update({"second_pass": "daily_quota_exhausted", "pass2": {"ok": False, "error": "מכסה יומית נגמרה: s1"}})
    wb, summary = save_and_load(tmp_path, [page])
    audit = wb["audit"]
    headers = [c.value for c in audit[1]]
    row = {h: c.value for h, c in zip(headers, next(audit.iter_rows(min_row=2)))}
    assert row["מעבר שני: סטטוס"] == "מכסה יומית נגמרה"
    assert summary["second_pass"] == {"מכסה יומית נגמרה": 1}


def test_build_rows_returns_cells_for_every_column():
    rows = build_rows(good_page())
    assert len(rows) == 1 and set(rows[0].cells) == {k for k, _ in COLUMNS} == set(export.FIELD_KEYS)


# ---------- עמודות טקסט: אפס מוביל לא נעלם ----------

def test_id_center_commitment_are_text_cells_and_keep_leading_zeros(tmp_path):
    page = good_page(1, id_numbers=[n("0011223344", "ת.ז.")],
                     numbers=[n("093937", "מרכזי"), n("00730578", "התחייבות מספר")])
    wb, _ = save_and_load(tmp_path, [page])          # הקובץ נכתב לדיסק ונקרא בחזרה
    ws = wb["תוצאות"]
    keys = [k for k, _ in COLUMNS]
    expect = {"id": "0011223344", "center": "093937", "commitment": "00730578"}
    for key, value in expect.items():
        c = ws.cell(row=2, column=keys.index(key) + 1)
        assert c.value == value and isinstance(c.value, str) and c.data_type == "s"
        assert c.number_format == "@"


def test_columns_are_the_seven_required_fields_in_the_brief_order():
    """דרישת ההגשה (CLAUDE.md סעיף 1): שבע עמודות, בסדר הזה בדיוק. אותו COLUMNS משמש גם את ייצוא
    הרשומות מהמסד (src/db_export.py), ולכן בדיקה אחת מכסה את שלושת הפורמטים."""
    assert [k for k, _ in COLUMNS] == ["first_name", "last_name", "id", "center", "date", "sender", "commitment"]
    assert [t for _, t in COLUMNS] == ["שם פרטי", "שם משפחה", 'ת"ז', "מרכזי", "תאריך קבלת דגימה", "גורם שולח", "מס' התחייבות"]
