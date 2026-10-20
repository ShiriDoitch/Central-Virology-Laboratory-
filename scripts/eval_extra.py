"""מדידה על סט הערכה נוסף וקטן: data/eval_extra/forms.pdf (6 עמודים), מול טבלת אמת ידנית ב-CSV.

מדידה בלבד. אין לשנות שום כלל, סף או קובץ הגדרות בעקבות הסט הזה (CLAUDE.md סעיף 9: כיוונון על חלק 1 בלבד).
הסט הזה אינו חלק 2 ואינו נוגע בו.

שימוש:
    python scripts/eval_extra.py                 # דוח בלבד, בלי אף קריאה למודל (מהתשובות השמורות ב-outputs/eval_extra)
    python scripts/eval_extra.py --run           # מריץ את הצינור על 6 העמודים (קורא ל-API! רק באישור)
    python scripts/eval_extra.py --run --pages 5,6   # מריץ עמודים נבחרים בלבד; הדוח תמיד על כל העמודים
    python scripts/eval_extra.py --truth=<path>  # קובץ אמת אחר
מטמון: במצב dev (ברירת המחדל) תשובות המודל נשמרות ב-data/cache, ולכן הרצה חוזרת של --run לא עולה קריאות.

טבלת האמת: xlsx או CSV, אותן עמודות. הסקריפט בוחר את הקובץ הראשון שקיים מתוך
data/eval_extra/truth.xlsx, truth.csv, truth_template.csv (ומדפיס במה השתמש); `--truth=<path>` גובר.
xlsx נוצר עם scripts/make_truth_xlsx.py ובו כל התאים בפורמט טקסט, כדי שאפסים מובילים בת"ז יישמרו.

פורמט טבלת האמת:
    page,first_name,last_name,id,center,receipt_date,sender,commitment
- **שורה לכל התחייבות.** עמוד עם שתי התחייבויות = שתי שורות עם אותו מספר עמוד; שם, ת"ז, מרכזי, תאריך וגורם שולח
  חוזרים בשתיהן, ורק ה-commitment שונה. השורות מקובצות לפי העמוד (TRUTH_ROW_FIELDS).
- **כמה מספרי מרכזי = תא אחד** מופרד ב-";" (למשל "944963 ; 944964"). אותו דבר לכל שדה רב-ערכי אחר.
- עמוד בלי התחייבות: שורה אחת עם commitment ריק.
- שדה שלא קיים בטופס: תא ריק. לא מנחשים.

איך נמדד (ראו build_report):
- ששת השדות האחרים נמדדים **לעמוד**: הצינור מעתיק את אותם פרטי מטופל לכל שורות העמוד (rules.expand_rows),
  ולכן השוואה לכל שורה הייתה סופרת את אותו תא פעמיים.
- **התחייבות נמדדת לכל שורה**: כל התחייבות באמת מותאמת לשורה שהמערכת הפיקה עם אותן ספרות (commitment_cells).
"""
import json
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src import extractor                                                      # noqa: E402
from src.assemble import page_rows                                             # noqa: E402
from src.evaluate import classify, compare_page, load_barcodes, load_run, norm_value   # noqa: E402
from src.pipeline import process_page, render_model_summary, sum_stats         # noqa: E402
from src.privacy import is_production, safe_write_text                         # noqa: E402
from src.rules import Decision, F_NO_COMMITMENT, apply_field_validity, split_multi     # noqa: E402
from src.senders import get_matcher                                            # noqa: E402
from run_step2_dates import BudgetReached, BudgetTracker                       # noqa: E402

PDF = ROOT / "data" / "eval_extra" / "forms.pdf"
TRUTH_XLSX = ROOT / "data" / "eval_extra" / "truth.xlsx"
TRUTH_CSV = ROOT / "data" / "eval_extra" / "truth.csv"
TRUTH_TEMPLATE = ROOT / "data" / "eval_extra" / "truth_template.csv"
TRUTH_CANDIDATES = [TRUTH_XLSX, TRUTH_CSV, TRUTH_TEMPLATE]      # הראשון שקיים הוא זה שנקרא
OUT = ROOT / "outputs" / "eval_extra"
PAGES = [1, 2, 3, 4, 5, 6]
ROTATED_PAGES = [5, 6]            # לפי המפתחת: צילומי מסך מסובבים של מסך מערכת

# תקרת קריאות קשיחה: בהגעה לתקרה הריצה נעצרת מיד ולא נשלחת אף בקשה נוספת.
# הועלתה ל-24 למודל החזק (20.10.2026) אחרי שריצת 6 העמודים נעצרה ב-14/14 כש-9 מ-14 הקריאות נבלעו ב-503 ובטיימאאוט:
# התקרה נגמרה מניסיונות חוזרים ולא מעבודה. אלה קבועים של סקריפט המדידה בלבד, לא מכסות המערכת (config/quotas.json).
LIMITS = {"gemini-3.5-flash-lite": 16, "gemini-3.1-flash-lite": 6, "gemini-3.6-flash": 24}

CSV_COLS = ["page", "first_name", "last_name", "id", "center", "receipt_date", "sender", "commitment"]
# עמודת ה-CSV -> שם השדה בקוד (receipt_date הוא "תאריך קבלת דגימה")
CSV_TO_FIELD = {"first_name": "first_name", "last_name": "last_name", "id": "id", "center": "center",
                "receipt_date": "date", "sender": "sender", "commitment": "commitment"}
MAIN_FIELDS = ["date", "center", "sender", "commitment"]
ID_FIELD = ["id"]
NAME_FIELDS = ["first_name", "last_name"]                # מדווחים בנפרד מת"ז (CLAUDE.md סעיף 16)
ALL_FIELDS = MAIN_FIELDS + ID_FIELD + NAME_FIELDS
# שדות שחוזרים בכל שורות העמוד (רק commitment שונה בין שורה לשורה)
TRUTH_ROW_FIELDS = [f for f in ALL_FIELDS if f != "commitment"]
# ארבע הקבוצות המבוקשות, מעל הקטגוריות של src/evaluate.py
BUCKETS = [("exact", ["confident_ok", "orange_ok"]), ("WRONG-GREEN", ["confident_wrong"]),
           ("wrong-orange", ["orange_wrong"]), ("empty", ["empty_ok", "empty_missed", "empty_flagged"])]
DETAIL = ["confident_ok", "confident_wrong", "orange_ok", "orange_wrong", "empty_ok", "empty_missed", "empty_flagged"]


def cell_text(value) -> str | None:
    """תא -> טקסט כפי שנכתב, או None אם ריק. אפסים מובילים נשמרים כשהתא הוא טקסט (פורמט "@").
    תא ש**אקסל** שמר כמספר איבד את האפסים המובילים עוד לפני שהקוד רואה אותו, ולכן תבנית ה-xlsx היא טקסט.
    מספר שלם: בלי ".0"; תאריך שאקסל המיר לתאריך אמיתי: dd.mm.yyyy (פורמט הפלט של המערכת)."""
    if value is None:
        return None
    if isinstance(value, str):
        return value.strip() or None
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip() or None


def _rows_from_csv(path: Path) -> tuple[list[str], list[dict]]:
    import csv
    with path.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        return list(reader.fieldnames or []), [{k: cell_text(v) for k, v in raw.items()} for raw in reader]


def _rows_from_xlsx(path: Path) -> tuple[list[str], list[dict]]:
    """קורא את הגיליון הראשון. כל התאים נקראים כטקסט (cell_text), בלי המרה למספר או לתאריך."""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = [cell_text(v) or "" for v in next(rows, ())]
        out = []
        for raw in rows:
            values = list(raw) + [None] * (len(header) - len(raw))
            out.append({name: cell_text(v) for name, v in zip(header, values) if name})
        return header, out
    finally:
        wb.close()


def load_truth(path: Path) -> dict[int, list[dict]]:
    """טבלת אמת (xlsx או CSV) -> {עמוד: [שורות אמת]} בפורמט ש-compare_page מצפה לו.
    שורה לכל התחייבות: **כמה שורות עם אותו מספר עמוד מקובצות יחד** (ולא דורסות זו את זו).
    מרכזי (וכל שדה רב-ערכי) נשאר כפי שנכתב, תא אחד מופרד ב-";", והפיצול נעשה בהשוואה (src.evaluate.compare_page).
    תא התחייבות עם כמה ערכים מופרדים ב-";" מתפצל גם הוא לשורה לכל התחייבות, כדי לסבול את שני הפורמטים."""
    header, rows = _rows_from_xlsx(path) if path.suffix.lower() in (".xlsx", ".xlsm") else _rows_from_csv(path)
    missing = [c for c in CSV_COLS if c not in header]
    if missing:
        raise ValueError(f"{path.name}: missing columns {missing}")
    out: dict[int, list[dict]] = {}
    for cells in rows:
        if not cells.get("page"):
            continue
        page = int(cells["page"])
        base = {field: cells.get(col) for col, field in CSV_TO_FIELD.items()}
        base["check"] = None                   # אין שדות מוחרגים בסט הזה
        commitments = split_multi(base["commitment"]) or [None]
        out.setdefault(page, []).extend({**base, "commitment": c} for c in commitments)
    return out


def row_disagreements(rows: list[dict]) -> list[str]:
    """שדות שאמורים לחזור זהים בכל שורות העמוד ובפועל נכתבו שונים (טעות הקלדה בטבלת האמת).
    ההשוואה משתמשת בשורה הראשונה, ולכן הבדל כזה מדווח כאזהרה."""
    return [f for f in TRUTH_ROW_FIELDS if len({(r[f] or "") for r in rows}) > 1]


def truth_filled(rows: list[dict]) -> bool:
    """שורת אמת שכל התאים בה ריקים = עדיין לא מולאה ביד."""
    return any(r[f] for r in rows for f in ALL_FIELDS)


def truth_path(argv: list[str], candidates: list[Path] | None = None) -> Path:
    """קובץ האמת: `--truth=<path>` גובר; אחרת הראשון שקיים מתוך truth.xlsx, truth.csv, truth_template.csv.
    אם אף אחד לא קיים: הראשון ברשימה (ההודעה למשתמשת תאמר שהוא חסר)."""
    for a in argv:
        if a.startswith("--truth="):
            return Path(a.split("=", 1)[1])
    options = candidates or TRUTH_CANDIDATES
    return next((p for p in options if p.exists()), options[0])


def parse_pages(argv: list[str], pages: list[int] | None = None) -> list[int]:
    """--pages 5,6 (או --pages=5,6, או טווח --pages 1-3): אילו עמודים **להריץ**. ברירת מחדל: כולם.
    הדוח נבנה תמיד על כל העמודים, מקובצי p<N>.json שקיימים, כך שעמודים מריצה קודמת לא נעלמים ממנו."""
    all_pages = pages or PAGES
    spec = None
    for i, a in enumerate(argv):
        if a.startswith("--pages="):
            spec = a.split("=", 1)[1]
        elif a == "--pages" and i + 1 < len(argv):
            spec = argv[i + 1]
    if not spec:
        return list(all_pages)
    out: list[int] = []
    for part in spec.split(","):
        part = part.strip()
        if "-" in part:
            a, b = part.split("-", 1)
            out += range(int(a), int(b) + 1)
        elif part:
            out.append(int(part))
    unknown = [p for p in out if p not in all_pages]
    if unknown:
        raise ValueError(f"--pages: עמודים שאינם בסט: {unknown} (הסט: {all_pages})")
    return sorted(dict.fromkeys(out))


def run_pages(tracker: BudgetTracker, pages: list[int] | None = None) -> tuple[str | None, list[dict]]:
    """מריץ את הצינור הרגיל על העמודים שנבחרו ושומר p<N>.json לכל עמוד. מחזיר (סיבת עצירה, סטטיסטיקות)."""
    OUT.mkdir(parents=True, exist_ok=True)
    stopped, stats = None, []
    try:
        for p in pages or PAGES:
            res = process_page(str(PDF), p, debug_dir=OUT)
            safe_write_text(OUT / f"p{p}.json", json.dumps(res, ensure_ascii=False))
            stats += [res["stats"]["pass1"], res["stats"]["pass2"]]
            rot = res.get("rotation") or {}
            print(f"page {p}: pass1_ok={bool(res['pass1'].get('ok'))} second_pass={res['second_pass']} "
                  f"pass2_model={res.get('pass2_model')} rotation_applied={rot.get('applied')} "
                  f"calls={dict(tracker.count)}", flush=True)
    except BudgetReached as e:
        stopped = f"call limit {e}"
        print(f"STOPPED: {stopped}", flush=True)
    return stopped, stats


def _pred_commitments(pred_rows: list[dict]) -> list[tuple[str, bool]]:
    """שורות המערכת שיש בהן מספר התחייבות: (ספרות מנורמלות, האם מסומן לבדיקה).
    F_NO_COMMITMENT אינו דגל לבדיקה אלא "לא נמצאה התחייבות" (כמו ב-src.evaluate.compare_page)."""
    return [(norm_value("commitment", r["commitment"]),
             any(f != F_NO_COMMITMENT for f in r["commitment_flags"]))
            for r in pred_rows if r["commitment"]]


def _empty_flagged(pred_rows: list[dict]) -> bool:
    """האם שורה בלי מספר התחייבות סומנה לבדיקה (ריק מסומן, ולא ריק שקט)."""
    return any(f != F_NO_COMMITMENT for r in pred_rows if not r["commitment"] for f in r["commitment_flags"])


def commitment_cells(truth_rows: list[dict], pred_rows: list[dict]) -> list[dict]:
    """התחייבות נמדדת **לכל שורה** (שורה = התחייבות אחת), ולא לכל עמוד.
    כל התחייבות באמת מותאמת לשורה שהמערכת הפיקה עם אותן ספרות:
    - התאמה: exact (כתום אם השורה מסומנת, אחרת ירוק).
    - התחייבות באמת בלי התאמה, ויש לה בת זוג שהמערכת הפיקה ולא קיימת באמת: קריאה שגויה (ירוק שגוי או כתום שגוי).
    - התחייבות באמת שהמערכת לא הפיקה בכלל: ריק (מסומן או שגיאה בלי סימון).
    - שורה שהמערכת הפיקה ואין לה אמת בכלל (שורה עודפת): קריאה שגויה.
    - אין התחייבות באמת ואין במערכת: תא ריק אחד לעמוד (נכון או מסומן)."""
    truth_values = [norm_value("commitment", r["commitment"]) for r in truth_rows if r["commitment"]]
    remaining = _pred_commitments(pred_rows)
    empty_flagged = _empty_flagged(pred_rows)
    out, unmatched = [], []
    for t in truth_values:
        hit = next((i for i, (v, _) in enumerate(remaining) if v == t), None)
        if hit is None:
            unmatched.append(t)
            continue
        _v, flagged = remaining.pop(hit)
        out.append({"truth": t, "pred": t, "category": "orange_ok" if flagged else "confident_ok"})
    for t in unmatched:
        if remaining:                                      # המערכת קראה מספר אחר במקום
            v, flagged = remaining.pop(0)
            out.append({"truth": t, "pred": v, "category": "orange_wrong" if flagged else "confident_wrong"})
        else:                                              # המערכת לא הפיקה את ההתחייבות הזאת
            out.append({"truth": t, "pred": "", "category": "empty_flagged" if empty_flagged else "empty_missed"})
    for v, flagged in remaining:                           # שורה עודפת שאין לה אמת
        out.append({"truth": "", "pred": v, "category": "orange_wrong" if flagged else "confident_wrong"})
    if not out:                                            # אין התחייבות באמת ולא במערכת
        out.append({"truth": "", "pred": "", "category": "empty_flagged" if empty_flagged else "empty_ok"})
    return out


def canonical_sender(value: str | None) -> str:
    """שם גורם שולח -> השם הקנוני מהרשימה הסגורה (src/senders.py), כדי שכינוי ושם קנוני ייחשבו אותו גורם
    ("וולפסון" = "מרכז רפואי וולפסון"). שם שאינו ברשימה: הטקסט המנורמל כפי שהוא, ולכן ממשיך להיספר כשונה."""
    if not value or not value.strip():
        return ""
    canon, _is_lab = get_matcher().match(value, trusted=True)
    return norm_value("sender", canon or value)


def sender_cell(truth_rows: list[dict], pred_rows: list[dict]) -> dict:
    """תא גורם שולח, בהשוואה קנונית (מחליף את התא של compare_page; src לא משתנה)."""
    dec: Decision = pred_rows[0]["sender"] if pred_rows else Decision()
    truth, pred = canonical_sender(truth_rows[0]["sender"]), canonical_sender(dec.value)
    return {"truth": truth, "pred": pred, "category": classify(truth == pred, bool(pred), dec.uncertain),
            "flags": list(dec.flags), "only_handwritten": False}


def cells(truth: dict[int, list[dict]], pages: list[int]) -> dict[int, dict | None]:
    """{עמוד: השוואה לכל שדה}. עמוד שנכשל או שאין לו ריצה: None.
    commitment_rows = ההשוואה לכל התחייבות בנפרד; היא שנספרת בדוח, במקום תא ההתחייבות של compare_page
    (שמשווה את כל מספרי העמוד כקבוצה אחת). גורם שולח מושווה קנונית (sender_cell).
    apply_field_validity מחושב מחדש על תשובת המעבר השני השמורה: field_ok נכתב לתוך p<N>.json בזמן הריצה,
    ולכן בלי החישוב מחדש הדוח היה מודד את כלל התוויות שהיה בתוקף אז ולא את זה שבקוד עכשיו."""
    results, regions = load_run(OUT, pages)
    barcodes = load_barcodes(OUT, pages)
    out: dict[int, dict | None] = {}
    for p in pages:
        if not results[p].get("ok"):
            out[p] = None
            continue
        r2 = regions.get(p)
        if r2 and r2.get("ok") and "patient" in r2:
            r2 = apply_field_validity(r2)
        rows = page_rows(results[p], regions=r2, barcodes=barcodes.get(p))
        out[p] = {**compare_page(truth[p], rows), "sender": sender_cell(truth[p], rows),
                  "commitment_rows": commitment_cells(truth[p], rows)}
    return out


def categories(cmp: dict[int, dict | None], field: str, pages: list[int]) -> Counter:
    """ספירת הקטגוריות לשדה. התחייבות: לכל שורה; שאר השדות: לכל עמוד."""
    if field == "commitment":
        return Counter(c["category"] for p in pages if cmp.get(p) for c in cmp[p]["commitment_rows"])
    return Counter(cmp[p][field]["category"] for p in pages if cmp.get(p))


def meta(pages: list[int]) -> dict[int, dict]:
    out = {}
    for p in pages:
        f = OUT / f"p{p}.json"
        out[p] = json.loads(f.read_text(encoding="utf-8")) if f.exists() else {}
    return out


def bucket_table(cmp: dict[int, dict | None], fields: list[str], pages: list[int]) -> list[str]:
    lines = ["| field | n | " + " | ".join(name for name, _ in BUCKETS) + " |", "|---|---|" + "---|" * len(BUCKETS)]
    for f in fields:
        c = categories(cmp, f, pages)
        n = sum(c[k] for k in DETAIL)
        unit = " (per commitment row)" if f == "commitment" else ""
        lines.append(f"| {f}{unit} | {n} | " + " | ".join(str(sum(c[k] for k in keys)) for _, keys in BUCKETS) + " |")
    return lines


def detail_table(cmp: dict[int, dict | None], pages: list[int]) -> list[str]:
    lines = ["| field | " + " | ".join(DETAIL) + " |", "|---|" + "---|" * len(DETAIL)]
    for f in ALL_FIELDS:
        c = categories(cmp, f, pages)
        lines.append(f"| {f} | " + " | ".join(str(c[k]) for k in DETAIL) + " |")
    return lines


def commitment_table(cmp: dict[int, dict | None], pages: list[int]) -> list[str]:
    """כל התחייבות בנפרד: האמת, מה שהמערכת הפיקה, והקטגוריה."""
    lines = ["| page | truth | pred | category |", "|---|---|---|---|"]
    for p in pages:
        if not cmp.get(p):
            lines.append(f"| {p} | (no run / page failed) | | |")
            continue
        for c in cmp[p]["commitment_rows"]:
            lines.append(f"| {p} | {c['truth']} | {c['pred']} | {c['category']} |")
    return lines


def rotation_table(info: dict[int, dict], pages: list[int]) -> list[str]:
    """זיהוי הכיוון לכל עמוד. "straight after rotation" = המעבר הראשון שנבחר דיווח rotation_degrees=0,
    כלומר המערכת מצאה את העמוד ישר (זו הראיה של המערכת עצמה, לא טבלת אמת)."""
    lines = ["| page | expected | local axis | model reported | rotation applied | attempts | unresolved | "
             "straight after rotation |", "|---|---|---|---|---|---|---|---|"]
    for p in pages:
        m = info.get(p) or {}
        rot = m.get("rotation") or {}
        p1 = m.get("pass1") or {}
        straight = p1.get("rotation_degrees") == "0" if p1.get("ok") else None
        expected = "rotated" if p in ROTATED_PAGES else "upright"
        lines.append(f"| {p} | {expected} | {rot.get('local_axis')} | {rot.get('reported')} | {rot.get('applied')} | "
                     f"{rot.get('attempts')} | {rot.get('unresolved')} | {straight} |")
    return lines


def iter_cells(cmp: dict[int, dict | None], pages: list[int]):
    """(עמוד, שדה, תא) לכל התאים הנמדדים. התחייבות: תא לכל שורה; שאר השדות: תא לכל עמוד."""
    for p in pages:
        if not cmp.get(p):
            continue
        for f in ALL_FIELDS:
            if f == "commitment":
                yield from ((p, f, c) for c in cmp[p]["commitment_rows"])
            else:
                yield p, f, cmp[p][f]


def build_report(truth: dict[int, list[dict]], pages: list[int], stopped: str | None, source: Path) -> str:
    cmp = cells(truth, pages)
    info = meta(pages)
    done = [p for p in pages if cmp.get(p)]
    lines = ["# Extra evaluation set (6 pages, measurement only)", "",
             f"Truth file: {source}  |  PDF: {PDF.name}  |  pages evaluated: {done or 'none'}",
             f"Truth rows per page: { {p: len(truth[p]) for p in pages} }", ""]
    for p in pages:
        bad = row_disagreements(truth[p])
        if bad:
            lines += [f"**Truth page {p}: these fields differ between its rows and should repeat: {bad}. "
                      "The first row is the one compared.**", ""]
    if stopped:
        lines += [f"**STOPPED: {stopped}. Report is partial.**", ""]
    missing = [p for p in pages if not cmp.get(p)]
    if missing:
        lines += [f"**No usable run for pages: {missing}** (run: python scripts/eval_extra.py --run)", ""]
    lines += ["## Main fields", ""] + bucket_table(cmp, MAIN_FIELDS, done)
    lines += ["", "## ID (reported separately, CLAUDE.md section 16)", ""] + bucket_table(cmp, ID_FIELD, done)
    lines += ["", "## Names (reported separately from ID)", ""] + bucket_table(cmp, NAME_FIELDS, done)
    lines += ["", "exact = correct value (green or orange)  |  WRONG-GREEN = wrong and not flagged (the metric to "
              "minimize)  |  wrong-orange = wrong but flagged  |  empty = no value (correct empty, missed error, or "
              "flagged empty)",
              "", "commitment is counted per commitment row (one truth row = one commitment, matched to the extracted "
              "row with the same digits); the other six fields are counted once per page, because the pipeline copies "
              "the same patient details into every row of a page.",
              "", "## Full category breakdown", ""] + detail_table(cmp, done)
    lines += ["", "## Commitment rows (one line per commitment)", ""] + commitment_table(cmp, pages)
    wrong_green = [(p, f, c["truth"], c["pred"]) for p, f, c in iter_cells(cmp, done)
                   if c["category"] == "confident_wrong"]
    lines += ["", "## WRONG GREEN (every case)", ""]
    lines += [f"- **page {p} {f}: truth {t!r} read {r!r}**" for p, f, t, r in wrong_green] or ["none"]
    lines += ["", "## Orientation detection", ""] + rotation_table(info, pages)
    lines += ["", "## Every cell that is not exact", "", "| page | field | truth | pred | category | flags |",
              "|---|---|---|---|---|---|"]
    for p in pages:
        if not cmp.get(p):
            lines.append(f"| {p} | (no run / page failed) | | | | |")
            continue
        for _p, f, c in iter_cells(cmp, [p]):
            if c["category"] not in ("confident_ok", "orange_ok"):
                lines.append(f"| {p} | {f} | {c['truth']} | {c['pred']} | {c['category']} | "
                             f"{', '.join(c.get('flags') or [])} |")
    lines += ["", "## Page status", "", "| page | pass 1 model | commitment form | quality | second pass | "
              "pass 2 model | Claude fallback |", "|---|---|---|---|---|---|---|"]
    for p in pages:
        m = info.get(p) or {}
        p1 = m.get("pass1") or {}
        claude = p1.get("provider") == "anthropic" or (m.get("pass2") or {}).get("provider") == "anthropic"
        lines.append(f"| {p} | {p1.get('model')} | {p1.get('is_commitment_form')} | {p1.get('page_quality')} | "
                     f"{m.get('second_pass')} | {m.get('pass2_model')} | {claude} |")
    return "\n".join(lines) + "\n"


def main(argv: list[str]) -> int:
    if is_production():
        print("PRIVACY_MODE=production: this script needs the p<N>.json dumps and the cache; run in dev mode.")
        return 1
    if not PDF.exists():
        print(f"missing PDF: {PDF}\nהעתיקי את הקובץ לשם (6 עמודים) ואז הריצי שוב.")
        return 1
    source = truth_path(argv)
    if not source.exists():
        print(f"missing truth file: {source}\n"
              f"אפשרויות שנבדקו: {', '.join(c.name for c in TRUTH_CANDIDATES)}. "
              "ליצירת תבנית ריקה: python scripts/make_truth_xlsx.py")
        return 1
    print(f"truth file: {source}")
    truth = load_truth(source)
    missing_pages = [p for p in PAGES if p not in truth]
    if missing_pages:
        print(f"{source.name}: missing rows for pages {missing_pages}")
        return 1
    empty = [p for p in PAGES if not truth_filled(truth[p])]
    if empty:
        print(f"אזהרה: בטבלת האמת {source.name} העמודים {empty} ריקים לגמרי (עוד לא מולאו ביד). "
              "כל תא ריק נמדד כאמת ריקה.")
    print(f"שורות אמת לכל עמוד: { {p: len(truth[p]) for p in PAGES} } "
          "(שורה לכל התחייבות; מרכזי מרובה בתא אחד מופרד ב-;)")
    for p in PAGES:
        bad = row_disagreements(truth[p])
        if bad:
            print(f"אזהרה: בעמוד {p} השדות {bad} שונים בין שורות העמוד, והם אמורים לחזור זהים. "
                  "ההשוואה משתמשת בשורה הראשונה.")
    if "--run" in argv:
        to_run = parse_pages(argv)
        print(f"מריצה את העמודים {to_run} (הדוח יכסה את כל {PAGES}, מקובצי p<N>.json שקיימים)")
        tracker = BudgetTracker(extractor.TRACKER, LIMITS)
        extractor.TRACKER = tracker
        stopped, stats = run_pages(tracker, to_run)
        print(render_model_summary(sum_stats(*stats)))
        print(f"calls this run: {dict(tracker.count)}  (hard limits: {LIMITS})")
    else:
        stopped = None
        print("report only: no model call was made. להרצה אמיתית: python scripts/eval_extra.py --run")
    text = build_report(truth, PAGES, stopped, source)
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "report.md").write_text(text, encoding="utf-8")
    print(text)
    print(f"saved: {OUT / 'report.md'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
