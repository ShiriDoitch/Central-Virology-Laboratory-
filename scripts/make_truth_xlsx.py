"""יוצר טבלת אמת ריקה כ-xlsx אמיתי (openpyxl) לסט ההערכה הנוסף: שורה לכל עמוד, כל התאים בפורמט טקסט.

למה טקסט: ת"ז עם אפסים מובילים (למשל 00112233) שנכתבת בתא רגיל נשמרת באקסל כמספר, והאפסים נעלמים
לפני שהקוד בכלל רואה את הקובץ. פורמט "@" (טקסט) על כל התאים ועל כל העמודה שומר את מה שהוקלד כפי שהוא.

שימוש:
    python scripts/make_truth_xlsx.py                      # data/eval_extra/truth.xlsx, 6 עמודים
    python scripts/make_truth_xlsx.py <path> [pages]
קובץ קיים לא נדרס (כדי לא לאבד מילוי ידני): צריך למחוק אותו קודם, או לתת שם אחר.
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUT = ROOT / "data" / "eval_extra" / "truth.xlsx"
DEFAULT_PAGES = 6
COLUMNS = ["page", "first_name", "last_name", "id", "center", "receipt_date", "sender", "commitment"]
TEXT_FORMAT = "@"               # פורמט "טקסט" של אקסל: שומר אפסים מובילים ולא הופך ערך למספר או לתאריך
WIDTHS = {"page": 6, "first_name": 14, "last_name": 14, "id": 14, "center": 20, "receipt_date": 14,
          "sender": 22, "commitment": 20}


def build_truth_xlsx(path: Path, pages: int = DEFAULT_PAGES, columns: list[str] | None = None) -> Path:
    """כותב xlsx עם שורת כותרת ושורה לכל עמוד (עמודת page מלאה, השאר ריק). כל התאים בפורמט טקסט."""
    columns = columns or COLUMNS
    wb = Workbook()
    ws = wb.active
    ws.title = "truth"
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.number_format = TEXT_FORMAT
    for p in range(1, pages + 1):
        ws.append([str(p)] + [None] * (len(columns) - 1))
    for row in ws.iter_rows(min_row=2, max_row=pages + 1, max_col=len(columns)):
        for cell in row:
            cell.number_format = TEXT_FORMAT
    for i, name in enumerate(columns, start=1):
        col = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[col].number_format = TEXT_FORMAT    # גם תאים שימולאו אחר כך בעמודה הזאת
        ws.column_dimensions[col].width = WIDTHS.get(name, 16)
    ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def main(argv: list[str]) -> int:
    out = Path(argv[0]) if argv else DEFAULT_OUT
    pages = int(argv[1]) if len(argv) > 1 else DEFAULT_PAGES
    if out.exists():
        print(f"{out} כבר קיים; לא נדרס. למחיקה ידנית ואז הרצה חוזרת, או שם קובץ אחר.")
        return 1
    build_truth_xlsx(out, pages)
    print(f"saved: {out} ({pages} rows, columns: {', '.join(COLUMNS)}, all cells text format)")
    print("שורה לכל התחייבות: עמוד עם שתי התחייבויות = שתי שורות עם אותו מספר עמוד. מרכזי מרובה: תא אחד מופרד ב-;")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
