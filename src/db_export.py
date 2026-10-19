"""ייצוא רשומות שמורות (מהמסד) ל-XLSX / CSV / JSON. מקבל רשומות מ-src/service.py (record_detail), בלי SQL.
הגנה מנוסחאות (= + - @) ב-XLSX וב-CSV; ב-JSON הערכים נשמרים כפי שהם. הקידוד של CSV הוא UTF-8 עם BOM כדי שאקסל יקרא עברית."""
import csv
import io
import json
from datetime import datetime

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

from .export import COLUMNS, CONFIRMED, FIELD_KEYS, GREEN, HEADER_FILL, ORANGE, TEXT_KEYS, YELLOW, protect
from .repo import FIELD_COLUMN

STATUS_HE = {"pending": "ממתינה", "ready": "מוכנה", "approved": "אושרה"}
EXTRA_COLUMNS = ["סטטוס", "קובץ", "עמוד", "מזהה רשומה"]
FORMATS = {
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
    "csv": ("text/csv", "csv"),
    "json": ("application/json", "json"),
}
FILL = {"green": GREEN, "orange": ORANGE, "yellow": YELLOW, "confirmed": CONFIRMED}


def _value(rec: dict, key: str) -> str:
    return (rec["center"] if key == "center" else rec[FIELD_COLUMN[key]]) or ""


def flat(rec: dict, filenames: dict[int, str]) -> list:
    return [*(_value(rec, k) for k in FIELD_KEYS), STATUS_HE[rec["status"]], filenames.get(rec["upload_id"], ""), rec["page_no"], rec["id"]]


def to_xlsx(recs: list[dict], filenames: dict[int, str]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "רשומות"
    ws.sheet_view.rightToLeft = True
    for i, title in enumerate([t for _, t in COLUMNS] + EXTRA_COLUMNS, start=1):
        c = ws.cell(row=1, column=i, value=title)
        c.fill = HEADER_FILL
    ws.freeze_panes = "A2"
    for r, rec in enumerate(recs, start=2):
        for i, v in enumerate(flat(rec, filenames), start=1):
            x = ws.cell(row=r, column=i, value=protect(v) if v != "" else None)
            if i <= len(FIELD_KEYS):
                key = FIELD_KEYS[i - 1]
                st = rec["fields"][key]
                if key in TEXT_KEYS:
                    x.number_format = "@"
                fill = FILL.get(st["color"])
                if fill is not None and v != "":
                    x.fill = fill
                elif st["color"] == "orange":
                    x.fill = ORANGE                                     # ריק מסומן לבדיקה
                notes = []
                if st["flags"]:
                    notes.append("דגלים: " + ", ".join(st["flags"]))
                if st["alternatives"]:
                    notes.append("חלופות: " + " | ".join(st["alternatives"]))
                if st["human_edited"]:
                    notes.append("תוקן ידנית" + (f"; הצעת המערכת: {st['machine_value']}" if st["machine_value"] else ""))
                elif st["cell_confirmed"]:
                    notes.append("אושר ידנית")
                if notes:
                    x.comment = Comment("\n".join(notes), "system", width=260, height=90)
    for letter, w in zip("ABCDEFGHIJK", [16, 18, 16, 22, 16, 30, 20, 12, 28, 8, 12]):
        ws.column_dimensions[letter].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_csv(recs: list[dict], filenames: dict[int, str]) -> bytes:
    out = io.StringIO()
    w = csv.writer(out, lineterminator="\r\n")
    w.writerow([t for _, t in COLUMNS] + EXTRA_COLUMNS)
    for rec in recs:
        w.writerow([protect(v) for v in flat(rec, filenames)])
    return out.getvalue().encode("utf-8-sig")


def to_json(recs: list[dict], filenames: dict[int, str]) -> bytes:
    items = []
    for rec in recs:
        item = {"id": rec["id"], "file": filenames.get(rec["upload_id"], ""), "page": rec["page_no"], "status": rec["status"],
                "reviewed_at": rec["reviewed_at"], "fields": {}}
        for k in FIELD_KEYS:
            st = rec["fields"][k]
            item["fields"][k] = {"value": _value(rec, k) or None, "color": st["color"], "flags": st["flags"],
                                 "alternatives": st["alternatives"], "machine_value": st["machine_value"],
                                 "human_edited": st["human_edited"], "confirmed": st["cell_confirmed"]}
        items.append(item)
    doc = {"exported_at": datetime.now().isoformat(timespec="seconds"), "count": len(items), "records": items}
    return json.dumps(doc, ensure_ascii=False, indent=2).encode("utf-8")


BUILDERS = {"xlsx": to_xlsx, "csv": to_csv, "json": to_json}
